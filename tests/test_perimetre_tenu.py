"""Le périmètre annoncé est-il tenu ?

Ces tests ne mesurent pas la qualité de détection. Ils vérifient une chose, et
une seule : **qu'aucune zone du fichier ne sort du traitement sans avoir été
lue**. Le témoin est toujours une adresse e-mail, motif que les règles
déterministes reconnaissent sans le modèle ; s'il survit, ce n'est pas que la
détection a hésité, c'est que personne n'a regardé.

C'est la contrepartie de l'encart « Périmètre du traitement » : l'utilisateur y
lit « traité » et transmet le fichier. Un élément annoncé sans test ici est une
promesse que rien ne garantit, et une lacune annoncée comme couverte est pire
qu'une lacune assumée.

Tous les tests tournent en mode dégradé (`NullNer`), donc sans le modèle : ils
sont rapides et ne dépendent d'aucun téléchargement.
"""
from datetime import datetime

import fitz
import openpyxl
import pytest

from anonymator.files import anonymize_file
from anonymator.files.pdf import pdf_io
from anonymator.ner import NullNer
from anonymator.referential import Referential
from tests import temoins_fixtures as fx

WHEN = datetime(2026, 9, 17, 10, 0, 0)


def _ref():
    return Referential.load_default()


def _run(path, out_dir):
    out_dir.mkdir(parents=True, exist_ok=True)
    return anonymize_file.anonymize_file(path, NullNer(), _ref(), out_dir, WHEN)


# --- Word ---------------------------------------------------------------

def test_docx_ne_laisse_aucun_temoin(tmp_path):
    src = fx.make_docx(tmp_path / "piege.docx")
    out = _run(src, tmp_path / "out").output_path
    assert fx.temoins_in_archive(out) == {}


@pytest.mark.parametrize("zone", [
    "corps", "entete", "pied", "insertion",
    "suppression",        # <w:delText> — révision suivie
    "champ",              # <w:instrText> — publipostage, lien
    "texte-lien",
    "cible-lien",         # word/_rels/document.xml.rels
    "customxml",
    "auteur-revision",    # attribut w:author
    "auteur", "titre",    # métadonnées
])
def test_docx_zone_par_zone(tmp_path, zone):
    """Un test par zone : un échec nomme la zone, pas « le docx »."""
    src = fx.make_docx(tmp_path / "piege.docx")
    out = _run(src, tmp_path / "out").output_path
    survivants = fx.temoins_in_archive(out)
    assert f"temoin.{zone}@example.com" not in survivants


def test_docx_reste_lisible_par_python_docx(tmp_path):
    """La réécriture de l'archive ne doit pas casser le document."""
    from docx import Document
    src = fx.make_docx(tmp_path / "piege.docx")
    out = _run(src, tmp_path / "out").output_path
    doc = Document(str(out))
    assert any("[EMAIL]" in p.text for p in doc.paragraphs)


# --- PowerPoint ---------------------------------------------------------

def test_pptx_ne_laisse_aucun_temoin(tmp_path):
    src = fx.make_pptx(tmp_path / "piege.pptx")
    out = _run(src, tmp_path / "out").output_path
    assert fx.temoins_in_archive(out) == {}


@pytest.mark.parametrize("zone", [
    "titre-slide", "zone-texte", "notes",
    "masque", "disposition",      # ppt/slideMasters, ppt/slideLayouts
    "auteur", "titre-doc",
])
def test_pptx_zone_par_zone(tmp_path, zone):
    src = fx.make_pptx(tmp_path / "piege.pptx")
    out = _run(src, tmp_path / "out").output_path
    assert f"temoin.{zone}@example.com" not in fx.temoins_in_archive(out)


# --- Classeur -----------------------------------------------------------

def test_xlsx_ne_laisse_aucun_temoin(tmp_path):
    src = fx.make_xlsx(tmp_path / "piege.xlsx")
    out = _run(src, tmp_path / "out").output_path
    assert fx.temoins_in_archive(out) == {}


@pytest.mark.parametrize("zone", [
    "cellule", "feuille-masquee",
    "intitule",              # ligne de titres
    "intervenant1",          # colonne à faible cardinalité
    "commentaire", "auteur-commentaire",
    "entete", "pied",        # en-tête / pied de feuille
    "nom-defini",
    "formule",               # littéral entre guillemets
    "auteur", "titre",
])
def test_xlsx_zone_par_zone(tmp_path, zone):
    src = fx.make_xlsx(tmp_path / "piege.xlsx")
    out = _run(src, tmp_path / "out").output_path
    assert f"temoin.{zone}@example.com" not in fx.temoins_in_archive(out)


def test_xlsx_reste_lisible_et_les_formules_survivent(tmp_path):
    """Masquer le littéral d'une formule ne doit pas détruire le calcul."""
    src = fx.make_xlsx(tmp_path / "piege.xlsx")
    out = _run(src, tmp_path / "out").output_path
    wb = openpyxl.load_workbook(out)
    # La formule reste une formule : seul le texte entre guillemets est réécrit.
    assert wb["Visible"]["C2"].value == '="[EMAIL]"'
    # La cellule, elle, est masquée par le plan de sa colonne : l'intitulé
    # « contact » la type PERSON, et toute la cellule est l'entité.
    assert wb["Visible"]["A2"].value == "[PERSONNE]"


def test_xlsx_cache_de_tableau_croise_retire(tmp_path):
    """Le cache d'un TCD est une copie littérale des lignes source : masquer la
    feuille n'y touche pas, il faut le retirer."""
    pytest.importorskip("openpyxl")
    src = _make_pivot_xlsx(tmp_path / "pivot.xlsx")
    out = _run(src, tmp_path / "out").output_path
    import zipfile
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        types = z.read("[Content_Types].xml").decode()
        rels = z.read("xl/_rels/workbook.xml.rels").decode()
    assert not [n for n in names if "pivotCache" in n or "pivotTable" in n]
    # Ni relation ni type de contenu orphelins : Excel proposerait de réparer.
    assert "pivot" not in types
    assert "pivotCache" not in rels
    assert fx.temoins_in_archive(out) == {}


def _make_pivot_xlsx(path):
    """Classeur minimal porteur d'un TCD et de son cache."""
    import shutil
    import zipfile
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Intervenant", "Montant"])
    ws.append(["temoin.pivot-a@example.com", 1])
    ws.append(["temoin.pivot-b@example.com", 2])
    wb.create_sheet("TCD")
    wb.save(path)

    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    pns = "http://schemas.openxmlformats.org/package/2006/relationships"
    cache_def = (
        f'<?xml version="1.0"?><pivotCacheDefinition xmlns="{ns}" xmlns:r="{rns}"'
        f' r:id="rId1" recordCount="2"><cacheSource type="worksheet">'
        f'<worksheetSource ref="A1:B3" sheet="Data"/></cacheSource>'
        f'<cacheFields count="2"><cacheField name="Intervenant" numFmtId="0">'
        f'<sharedItems count="2"><s v="temoin.pivot-a@example.com"/>'
        f'<s v="temoin.pivot-b@example.com"/></sharedItems></cacheField>'
        f'<cacheField name="Montant" numFmtId="0"><sharedItems '
        f'containsSemiMixedTypes="0" containsString="0" containsNumber="1" '
        f'minValue="1" maxValue="2"/></cacheField></cacheFields>'
        f'</pivotCacheDefinition>')
    cache_rec = (f'<?xml version="1.0"?><pivotCacheRecords xmlns="{ns}" count="2">'
                 f'<r><x v="0"/><n v="1"/></r><r><x v="1"/><n v="2"/></r>'
                 f'</pivotCacheRecords>')
    table = (
        f'<?xml version="1.0"?><pivotTableDefinition xmlns="{ns}" name="TCD1"'
        f' cacheId="1" dataCaption="Valeurs" updatedVersion="6"'
        f' minRefreshableVersion="3" createdVersion="6" indent="0" outline="1"'
        f' outlineData="1" multipleFieldFilters="0">'
        f'<location ref="A1:B4" firstHeaderRow="1" firstDataRow="1" firstDataCol="0"/>'
        f'<pivotFields count="2"><pivotField axis="axisRow" showAll="0">'
        f'<items count="3"><item x="0"/><item x="1"/><item t="default"/></items>'
        f'</pivotField><pivotField dataField="1" showAll="0"/></pivotFields>'
        f'<rowFields count="1"><field x="0"/></rowFields>'
        f'<rowItems count="3"><i><x/></i><i><x v="1"/></i><i t="grand"><x/></i></rowItems>'
        f'<colItems count="1"><i/></colItems>'
        f'<dataFields count="1"><dataField name="Somme" fld="1" baseField="0"'
        f' baseItem="0"/></dataFields></pivotTableDefinition>')

    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", (
                    '<Override PartName="/xl/pivotCache/pivotCacheDefinition1.xml"'
                    ' ContentType="application/vnd.openxmlformats-officedocument'
                    '.spreadsheetml.pivotCacheDefinition+xml"/>'
                    '<Override PartName="/xl/pivotCache/pivotCacheRecords1.xml"'
                    ' ContentType="application/vnd.openxmlformats-officedocument'
                    '.spreadsheetml.pivotCacheRecords+xml"/>'
                    '<Override PartName="/xl/pivotTables/pivotTable1.xml"'
                    ' ContentType="application/vnd.openxmlformats-officedocument'
                    '.spreadsheetml.pivotTable+xml"/></Types>').encode())
            if item.filename == "xl/workbook.xml":
                data = data.replace(
                    b'<workbook xmlns=',
                    f'<workbook xmlns:r="{rns}" xmlns='.encode(), 1)
                data = data.replace(
                    b"</workbook>",
                    b'<pivotCaches><pivotCache cacheId="1" r:id="rIdPC1"/>'
                    b'</pivotCaches></workbook>')
            if item.filename == "xl/_rels/workbook.xml.rels":
                data = data.replace(b"</Relationships>", (
                    f'<Relationship Id="rIdPC1" Type="{rns}/pivotCacheDefinition"'
                    f' Target="pivotCache/pivotCacheDefinition1.xml"/>'
                    f'</Relationships>').encode())
            zo.writestr(item, data)
        zo.writestr("xl/pivotCache/pivotCacheDefinition1.xml", cache_def)
        zo.writestr("xl/pivotCache/pivotCacheRecords1.xml", cache_rec)
        zo.writestr("xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels",
                    f'<?xml version="1.0"?><Relationships xmlns="{pns}">'
                    f'<Relationship Id="rId1" Type="{rns}/pivotCacheRecords"'
                    f' Target="pivotCacheRecords1.xml"/></Relationships>')
        zo.writestr("xl/pivotTables/pivotTable1.xml", table)
        zo.writestr("xl/pivotTables/_rels/pivotTable1.xml.rels",
                    f'<?xml version="1.0"?><Relationships xmlns="{pns}">'
                    f'<Relationship Id="rId1" Type="{rns}/pivotCacheDefinition"'
                    f' Target="../pivotCache/pivotCacheDefinition1.xml"/>'
                    f'</Relationships>')
        zo.writestr("xl/worksheets/_rels/sheet2.xml.rels",
                    f'<?xml version="1.0"?><Relationships xmlns="{pns}">'
                    f'<Relationship Id="rId1" Type="{rns}/pivotTable"'
                    f' Target="../pivotTables/pivotTable1.xml"/></Relationships>')
    shutil.move(str(tmp), str(path))
    return path


# --- CSV ----------------------------------------------------------------

@pytest.mark.parametrize("zone", [
    "intitule-a", "intitule-b",   # identités en ligne de titres
    "intervenant1",               # colonne à faible cardinalité
])
def test_csv_zone_par_zone(tmp_path, zone):
    """Le tableau croisé est la structure la plus banale d'un export métier, et
    c'était celle qui sortait intégralement en clair avec « aucune détection »."""
    src = fx.make_csv(tmp_path / "piege.csv")
    out_dir = tmp_path / "out"
    out = _run(src, out_dir).output_path
    assert f"temoin.{zone}@example.com" not in out.read_text(encoding="utf-8")


def test_csv_conserve_les_intitules_de_colonnes_ordinaires(tmp_path):
    """La contrepartie : un intitulé reconnu comme nom de colonne reste intact.
    Le masquer casserait le fichier sans rien protéger."""
    src = fx.make_csv(tmp_path / "piege.csv")
    out = _run(src, tmp_path / "out").output_path
    first = out.read_text(encoding="utf-8").splitlines()[0]
    assert first.split(";")[0] == "mois"
    assert first.split(";")[4] == "acte"


# --- PDF ----------------------------------------------------------------

def _redact_all(src, out_dir):
    """Caviarde tout ce que le scan détecte, comme le ferait l'utilisateur qui
    coche tout dans la revue."""
    ref = _ref()
    out_dir.mkdir(parents=True, exist_ok=True)
    scans = pdf_io.scan_pdf(src, NullNer(), ref)
    rects = {}
    for page in scans:
        boxes = [w.rect for e in page.entities for w in page.words
                 if not (w.char_end <= e.start or w.char_start >= e.end)]
        if boxes:
            rects[page.page_index] = boxes
    return pdf_io.anonymize_pdf_redact(src, rects, out_dir, WHEN,
                                       ner=NullNer(), ref=ref)


def test_pdf_ne_laisse_aucun_temoin_dans_les_octets(tmp_path):
    src = fx.make_pdf(tmp_path / "piege.pdf")
    out = _redact_all(src, tmp_path / "out")
    assert fx.temoins_in(out.read_bytes()) == set()


def test_pdf_champ_de_formulaire_vide(tmp_path):
    """Le pire cas : la valeur est extraite, donc détectée, affichée à la revue
    et cochée — et `apply_redactions` ne la touche pas, elle vit hors du flux
    de page. Elle ressortait d'un document présenté comme caviardé."""
    src = fx.make_pdf(tmp_path / "piege.pdf")
    out = _redact_all(src, tmp_path / "out")
    doc = fitz.open(str(out))
    assert [w.field_value for p in doc for w in (p.widgets() or [])] == [""]


def test_pdf_annotations_masquees_et_auteur_efface(tmp_path):
    src = fx.make_pdf(tmp_path / "piege.pdf")
    out = _redact_all(src, tmp_path / "out")
    doc = fitz.open(str(out))
    infos = [a.info for p in doc for a in (p.annots() or [])]
    assert infos, "les annotations ne doivent pas disparaître, seulement être masquées"
    for info in infos:
        assert "temoin" not in (info.get("content") or "")
        assert "temoin" not in (info.get("title") or "")


def test_pdf_signets_masques_et_pieces_jointes_retirees(tmp_path):
    src = fx.make_pdf(tmp_path / "piege.pdf")
    out = _redact_all(src, tmp_path / "out")
    doc = fitz.open(str(out))
    assert doc.embfile_count() == 0
    assert all("temoin" not in (entry[1] or "") for entry in doc.get_toc())


def test_pdf_texte_de_page_reellement_detruit(tmp_path):
    """Acquis à préserver : la rédaction n'est pas un rectangle dessiné."""
    src = fx.make_pdf(tmp_path / "piege.pdf")
    out = _redact_all(src, tmp_path / "out")
    doc = fitz.open(str(out))
    assert "temoin.corps@example.com" not in "".join(p.get_text() for p in doc)


def test_pdf_metadonnees_et_xmp_vides(tmp_path):
    src = fx.make_pdf(tmp_path / "piege.pdf")
    out = _redact_all(src, tmp_path / "out")
    doc = fitz.open(str(out))
    assert not doc.metadata.get("author")
    assert not doc.metadata.get("title")
