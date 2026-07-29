"""Le jeu de démonstration est le cas de référence du chantier : colonnes
d'identités, nomenclatures et mesures dans un même fichier, en CSV et en XLSX."""
from pathlib import Path

import openpyxl
import pytest

from anonymator.referential import Referential
from anonymator.ner import NullNer
from anonymator.files import csv_io, xlsx_io
from anonymator.files.anonymize_file import csv_column_plans, scan_csv
from anonymator.files.columns import SKIP, TYPED, classify_columns
from anonymator.core.file_review_session import FileReviewSession
from anonymator.core.xlsx_review_session import XlsxReviewSession
from anonymator.core.tabular_review_session import CLEAR, MASK

CSV = Path("exemples/clients_demo.csv")
XLSX = Path("exemples/clients_demo.xlsx")

pytestmark = pytest.mark.skipif(not CSV.exists(),
                                reason="jeu de démonstration absent")


def _csv_session():
    doc = csv_io.read_csv(CSV)
    doc.has_header = True
    ref = Referential.load_default()
    plans = classify_columns(doc.rows, doc.has_header)
    scan_plans = csv_column_plans(doc)
    scanned = scan_csv(doc, NullNer(), ref, scan_plans)
    return FileReviewSession(doc, scanned, ref, set(scan_plans), plans)


def test_demo_csv_plan_separates_identities_measures_and_nomenclatures():
    s = _csv_session()
    by_name = {h: c for c, h in enumerate(s.doc.rows[0])}
    assert s.plans[by_name["contact_nom"]].policy == TYPED
    assert s.plans[by_name["telephone"]].policy == TYPED
    assert s.plans[by_name["ca_2025"]].policy == SKIP
    assert s.plans[by_name["statut"]].policy == SKIP


def test_demo_csv_column_override_masks_a_skipped_column(tmp_path):
    """Forcer une nomenclature écartée par le plan : 120 cellules masquées."""
    s = _csv_session()
    by_name = {h: c for c, h in enumerate(s.doc.rows[0])}
    col = by_name["secteur"]
    # `raison_sociale` est déjà typée ORG par son en-tête : le forçage de
    # `secteur` s'ajoute aux 120 occurrences existantes, il ne les remplace pas.
    assert s.count_retained("ORG") == 120
    s.set_column_override(col, MASK, "ORG")
    assert s.count_retained("ORG") == 240
    out = tmp_path / "demo.csv"
    s.apply_and_save(out)
    lines = out.read_bytes().decode(s.doc.encoding).splitlines()
    assert lines[0].split(";")[col] == "secteur"          # titres intacts
    assert lines[1].split(";")[col] == "[ORG]"


def test_demo_csv_column_override_frees_a_typed_column(tmp_path):
    s = _csv_session()
    by_name = {h: c for c, h in enumerate(s.doc.rows[0])}
    s.set_column_override(by_name["email"], CLEAR)
    out = tmp_path / "demo.csv"
    s.apply_and_save(out)
    text = out.read_bytes().decode(s.doc.encoding)
    assert "sandra.leclerc@bureau-sablons.net" in text
    assert "[PERSONNE]" in text                            # le reste est traité


@pytest.mark.skipif(not XLSX.exists(), reason="conversion xlsx absente")
def test_demo_xlsx_reviews_like_the_csv(tmp_path):
    ref = Referential.load_default()
    res = xlsx_io.scan_workbook(XLSX, NullNer(), ref)
    s = XlsxReviewSession(res, ref)
    assert res.has_header["Clients"] is True
    assert "PERSON" in s.types() and "EMAIL" in s.types()
    by_name = {h: c for c, h in enumerate(res.matrices["Clients"][0])}
    assert s.plans[("Clients", by_name["ca_2025"])].policy == SKIP
    s.set_column_override(("Clients", by_name["secteur"]), MASK, "ORG")
    out = tmp_path / "demo.xlsx"
    s.apply_and_save(out)
    ws = openpyxl.load_workbook(out)["Clients"]
    col = by_name["secteur"] + 1
    assert ws.cell(row=1, column=col).value == "secteur"
    assert ws.cell(row=2, column=col).value == "[ORG]"
    assert ws.cell(row=2, column=by_name["contact_nom"] + 1).value == "[PERSONNE]"
    assert ws.cell(row=2, column=by_name["ca_2025"] + 1).value == 5230650
