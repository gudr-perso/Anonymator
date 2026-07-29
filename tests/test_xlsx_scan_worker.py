import openpyxl

from anonymator.referential import Referential
from anonymator.ner import FakeNer
from anonymator.ui.model_loader import ModelLoader
from anonymator.ui.xlsx_scan_worker import XlsxScanWorker


def _book(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clients"
    ws["A1"] = "contact_nom"; ws["B1"] = "ca_2025"
    for i in range(4):
        ws.cell(row=i + 2, column=1, value=["Leclerc", "Berger"][i % 2])
        ws.cell(row=i + 2, column=2, value=100 + i)
    wb.save(path)
    return path


def test_worker_emits_scan_result(qtbot, tmp_path):
    src = _book(tmp_path / "b.xlsx")
    worker = XlsxScanWorker(src, ModelLoader(FakeNer({})),
                            Referential.load_default())
    with qtbot.waitSignal(worker.scan_finished, timeout=10000) as blocker:
        worker.start()
    worker.wait()
    res = blocker.args[0]
    assert res.sheets == ["Clients"]
    assert ("Clients", 1, 0) in res.scanned


def test_worker_passes_header_overrides(qtbot, tmp_path):
    src = _book(tmp_path / "b.xlsx")
    worker = XlsxScanWorker(src, ModelLoader(FakeNer({})),
                            Referential.load_default(),
                            header_overrides={"Clients": False})
    with qtbot.waitSignal(worker.scan_finished, timeout=10000) as blocker:
        worker.start()
    worker.wait()
    assert blocker.args[0].has_header["Clients"] is False


def test_worker_reports_a_load_failure(qtbot, tmp_path):
    """Le détecteur est construit DANS le thread : un échec remonte via
    `error` au lieu d'exploser en silence sur le thread principal."""
    src = _book(tmp_path / "b.xlsx")

    class BoomLoader:
        def get(self):
            raise RuntimeError("échec chargement modèle")

    worker = XlsxScanWorker(src, BoomLoader(), Referential.load_default())
    with qtbot.waitSignal(worker.error, timeout=10000) as blocker:
        worker.start()
    worker.wait()
    assert "échec chargement modèle" in blocker.args[0]
