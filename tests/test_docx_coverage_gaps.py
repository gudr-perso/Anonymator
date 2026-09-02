"""Conteneurs docx que l'API python-docx laissait dans l'ombre.

Trois trous constatés sur des documents réels, tous refermés par le parcours
XML de `docx_io` : le texte d'un lien hypertexte, celui d'un contrôle de
contenu, et les cellules fusionnées (visitées deux fois, donc masquées deux
fois — la seconde passe corrompait le texte).
"""
from datetime import datetime

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from anonymator.ner import NullNer
from anonymator.referential import Referential
from anonymator.files.ooxml import docx_io

_WHEN = datetime(2026, 9, 2, 10, 0, 0)


def _anonymize(tmp_path, doc):
    src = tmp_path / "src.docx"
    doc.save(str(src))
    out, report = docx_io.anonymize_document(
        src, NullNer(), Referential.load_default(), tmp_path, when=_WHEN)
    return Document(str(out)), report


def _add_hyperlink_run(paragraph, text: str) -> None:
    """Un <w:r> enveloppé dans <w:hyperlink>, comme Word en produit dès qu'une
    adresse e-mail est saisie."""
    link = etree.SubElement(paragraph._p, qn("w:hyperlink"))
    run = etree.SubElement(link, qn("w:r"))
    etree.SubElement(run, qn("w:t")).text = text


def _add_content_control(doc, text: str) -> None:
    """Un paragraphe dans <w:sdt>/<w:sdtContent> — la forme d'un champ de
    formulaire ou d'un modèle Word."""
    sdt = etree.SubElement(doc.element.body, qn("w:sdt"))
    content = etree.SubElement(sdt, qn("w:sdtContent"))
    p = etree.SubElement(content, qn("w:p"))
    run = etree.SubElement(p, qn("w:r"))
    etree.SubElement(run, qn("w:t")).text = text


def test_hyperlink_text_is_masked(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Écrire à ")
    _add_hyperlink_run(p, "jean.dupont@exemple.fr")
    out, report = _anonymize(tmp_path, doc)
    body = out.element.body.xml
    assert "jean.dupont@exemple.fr" not in body
    assert "[EMAIL]" in body
    assert any(r["original"] == "jean.dupont@exemple.fr" for r in report.to_rows())


def test_content_control_paragraph_is_masked(tmp_path):
    doc = Document()
    _add_content_control(doc, "IBAN FR7630006000011234567890189")
    out, _report = _anonymize(tmp_path, doc)
    body = out.element.body.xml
    assert "FR7630006000011234567890189" not in body
    assert "[IBAN]" in body


def test_merged_cell_is_masked_once(tmp_path):
    """`row.cells` rend le même objet pour chaque colonne d'une fusion : la
    cellule était masquée deux fois et le second passage réappliquait les
    offsets d'origine sur un texte déjà masqué."""
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "Mail jean@ex.com puis paul@ex.com fin"
    out, report = _anonymize(tmp_path, doc)
    assert out.tables[0].cell(0, 0).text == "Mail [EMAIL] puis [EMAIL] fin"
    occurrences = {r["original"]: r["occurrences"] for r in report.to_rows()}
    assert occurrences["jean@ex.com"] == 1


def test_first_page_header_is_masked(tmp_path):
    """Une donnée nominative peut ne figurer que sur l'en-tête de première
    page, distinct de l'en-tête normal."""
    doc = Document()
    doc.add_paragraph("Corps")
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = "Dossier de jean@ex.com"
    out, _report = _anonymize(tmp_path, doc)
    assert "[EMAIL]" in out.sections[0].first_page_header.paragraphs[0].text
