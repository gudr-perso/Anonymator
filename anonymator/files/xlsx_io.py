from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import openpyxl

from anonymator.model import Entity
from anonymator.ner import NerDetector
from anonymator.referential import Referential
from anonymator.pipeline import detect
from anonymator.anonymize import apply_masking
from anonymator.dedup import detect_unique
from anonymator.report.audit import AuditReport
from anonymator.output_naming import anonymized_path
from anonymator.files import xlsx_parts
from anonymator.files.ooxml import scan as unit_scan
from anonymator.files.columns import (
    ColumnPlan, classify_columns, column_detector, header_cells_to_scan,
    looks_like_header_row)


# Propriétés de `docProps/core.xml` portant une identité ou un contenu, et
# libellé d'audit associé. Les horodatages (created/modified) restent : ils ne
# désignent personne, et les retirer casserait des outils qui s'y fient.
_CORE_PROPS = [
    ("creator", "Auteur"),
    ("lastModifiedBy", "Dernier modifié par"),
    ("title", "Titre"),
    ("subject", "Sujet"),
    ("keywords", "Mots-clés"),
    ("description", "Commentaires"),
    ("category", "Catégorie"),
    ("identifier", "Identifiant"),
]


def _is_formula(cell) -> bool:
    return cell.data_type == "f" or (isinstance(cell.value, str)
                                     and cell.value.startswith("="))


def _cell_text(cell) -> str:
    """Représentation texte d'une cellule pour la classification.

    Les formules comptent pour vides : leur source n'est pas une donnée, et
    elles ne doivent jamais peser sur le profil d'une colonne."""
    if cell.value is None or _is_formula(cell):
        return ""
    return str(cell.value)


def sheet_has_header(ws) -> bool:
    """Première ligne = en-têtes ?

    Contrairement au CSV, le classeur porte le type de chaque cellule : une
    ligne de titres est faite de texte, au-dessus d'au moins une colonne qui
    ne l'est pas. C'est un fait lu dans le fichier, pas une statistique — il
    reste donc le signal décisif, et un refus de sa part n'est jamais annulé.

    Il a un angle mort, et un seul : une feuille dont *toutes* les colonnes
    sont textuelles. Là, et là seulement, on regarde si la ligne 1 est faite
    de noms de colonnes connus (cf. columns.looks_like_header_row). Ce repli
    n'invente rien sur la forme des chaînes : il interroge le lexique qui
    pilote déjà le typage. Sans donnée en dessous, ou sans vocabulaire
    reconnu, la ligne est traitée comme une donnée — choix prudent : elle est
    analysée plutôt qu'ignorée."""
    rows = list(ws.iter_rows())
    if len(rows) < 2:
        return False
    first = [c for c in rows[0] if c.value is not None]
    if not first or not all(isinstance(c.value, str) and not _is_formula(c)
                            for c in first):
        return False
    if any(cell.value is not None and not isinstance(cell.value, str)
           for row in rows[1:] for cell in row):
        return True
    return looks_like_header_row([c.value for c in first])


def _sheet_matrix(ws) -> list[list[str]]:
    return [[_cell_text(c) for c in row] for row in ws.iter_rows()]


@dataclass
class XlsxScanResult:
    """Tout ce qu'une revue doit connaître d'un classeur, sans rien y écrire.

    Le classeur openpyxl reste ouvert : c'est lui qu'on masquera à la fin, ce
    qui préserve la mise en forme et les formules."""
    workbook: object
    sheets: list[str]
    matrices: dict[str, list[list[str]]]
    plans: dict[str, dict[int, ColumnPlan]]
    has_header: dict[str, bool]
    scanned: dict[tuple[str, int, int], list[Entity]] = field(default_factory=dict)
    # Valeurs d'origine des cellules déjà masquées, pour pouvoir les rendre
    # avant un nouveau passage (cf. apply_workbook).
    originals: dict[tuple[str, int, int], object] = field(default_factory=dict)
    # Détecteur ayant servi au scan, conservé pour la passe hors grille
    # (commentaires, en-têtes, formules) exécutée à l'enregistrement.
    ner: object = None


def scan_workbook(path: Path, ner: NerDetector, ref: Referential,
                  header_overrides: dict[str, bool] | None = None) -> XlsxScanResult:
    """Lit le classeur, classe ses colonnes et détecte les entités, sans rien
    écrire. Séparer le scan de l'application est ce qui rend la revue possible :
    l'utilisateur tranche entre les deux (cf. files/ooxml/scan.py).

    Le plan de traitement est décidé une fois par colonne et par feuille (cf.
    columns.py) : une colonne typée est masquée en entier sans passer par le
    NER, une nomenclature ou une mesure reste intacte.

    Indices de `scanned` : ceux de la matrice, donc décalés de 1 par rapport
    aux coordonnées openpyxl (ligne 1 du classeur = ligne 0 de la matrice)."""
    wb = openpyxl.load_workbook(path)
    overrides = header_overrides or {}
    sheets: list[str] = []
    matrices: dict[str, list[list[str]]] = {}
    plans: dict[str, dict[int, ColumnPlan]] = {}
    headers: dict[str, bool] = {}
    scanned: dict[tuple[str, int, int], list[Entity]] = {}
    for ws in wb.worksheets:
        title = ws.title
        sheets.append(title)
        matrix = _sheet_matrix(ws)
        matrices[title] = matrix
        if not matrix:
            plans[title] = {}
            headers[title] = False
            continue
        has_header = overrides.get(title, sheet_has_header(ws))
        headers[title] = has_header
        sheet_plans = classify_columns(matrix, has_header)
        plans[title] = sheet_plans
        start = 1 if has_header else 0
        for col, plan in sheet_plans.items():
            detector = column_detector(plan, ner, ref)
            rows = [r for r in range(start, len(matrix))
                    if col < len(matrix[r]) and matrix[r][col]]
            cache = detect_unique([matrix[r][col] for r in rows], detector)
            for r in rows:
                ents = cache.get(matrix[r][col], [])
                if ents:
                    scanned[(title, r, col)] = ents
        # Ligne de titres : détecteur complet, pas celui de la colonne.
        for col in header_cells_to_scan(matrix, has_header):
            ents = detect(matrix[0][col], ner, ref)
            if ents:
                scanned[(title, 0, col)] = ents
    return XlsxScanResult(wb, sheets, matrices, plans, headers, scanned, ner=ner)


def restore_cells(result: XlsxScanResult) -> None:
    """Rend leur valeur d'origine aux cellules déjà masquées.

    Le classeur openpyxl est modifié en place et reste ouvert entre deux
    enregistrements : sans cette remise à zéro, un second passage réappliquait
    les offsets d'origine sur un texte déjà masqué — « [PERSONNE] » devenait
    « [PERSONNE]NNE] », silencieusement."""
    for (title, r, c), value in result.originals.items():
        result.workbook[title].cell(row=r + 1, column=c + 1).value = value
    result.originals.clear()


def apply_workbook(result: XlsxScanResult,
                   retained: dict[tuple[str, int, int], list[Entity]],
                   ref: Referential,
                   report: AuditReport | None = None) -> AuditReport:
    """Écrit les entités retenues dans les cellules du classeur.

    Rejouable : les cellules masquées lors d'un passage précédent sont d'abord
    rendues à leur valeur d'origine.

    Une cellule de formule n'est jamais réécrite : sa source n'est pas une
    donnée, et l'écraser détruirait le calcul."""
    report = report if report is not None else AuditReport()
    restore_cells(result)
    for key, ents in retained.items():
        if not ents:
            continue
        title, r, c = key
        cell = result.workbook[title].cell(row=r + 1, column=c + 1)
        if _is_formula(cell):
            continue
        value = _cell_text(cell)
        location = f"{title}!{cell.coordinate}"
        for e in ents:
            report.add(e.type, e.value, ref.tag_for(e.type), location)
        result.originals[key] = cell.value
        cell.value = apply_masking(value, ents, ref)
    return report


def purge_metadata(workbook, report: AuditReport) -> AuditReport:
    """Vide les propriétés de document du classeur (auteur, titre, sujet…).

    openpyxl réécrit `docProps/core.xml` depuis `workbook.properties` : sans
    cette purge, le classeur anonymisé sortait avec le nom de son auteur et son
    titre d'origine — « Paie 2026 » en dit parfois plus que son contenu. Les
    autres formats les purgeaient déjà (docx/pptx via `ooxml.metadata`, PDF via
    `pdf.redact`) ; le classeur était le seul trou.

    Rejouable : après un premier passage il ne reste rien à signaler."""
    props = workbook.properties
    for attr, label in _CORE_PROPS:
        value = getattr(props, attr, None)
        if value is not None and str(value).strip():
            report.add("META", str(value), "", f"Métadonnées / {label}")
            setattr(props, attr, None)
    return report


def postprocess_workbook(workbook, ner, ref: Referential,
                         report: AuditReport) -> AuditReport:
    """Traite ce que la grille ne montre pas, juste avant l'enregistrement.

    Commentaires de cellule et leur auteur, en-têtes et pieds de page, noms
    définis, littéraux de formule : autant de porteurs de texte que
    `scan_workbook` ne voit pas, puisqu'il ne lit que `cell.value`. Ils suivent
    ici le chemin des commentaires d'un .docx — détection, puis masquage des
    seules entités confirmées, sans passer par la revue, une donnée invisible à
    l'écran ne pouvant pas être arbitrée à l'écran.

    Les caches de tableau croisé sont retirés : ils contiennent une copie non
    masquée des lignes source.

    Rejouable : un second enregistrement retrouve un texte déjà masqué, sur
    lequel plus rien ne se détecte."""
    units = xlsx_parts.hidden_text_units(workbook)
    if units:
        retained = unit_scan.confirmed_only(unit_scan.scan_units(units, ner, ref))
        if retained:
            unit_scan.apply_units(units, retained, ref, report)
    removed = xlsx_parts.drop_pivot_caches(workbook)
    for _ in range(removed):
        report.add("META", "tableau croisé dynamique", "",
                   "Cache de tableau croisé retiré")
    return report


def anonymize_workbook(path: Path, ner: NerDetector, ref: Referential,
                       output_dir: Path, when: datetime) -> tuple[Path, AuditReport]:
    """Chemin direct, sans revue : scan puis application immédiate.

    Sans revue, il n'y a pas d'opt-in : on filtre les entités non confirmées
    (format plausible mais contrôle de clé KO) plutôt que de masquer une
    fausse détection. Même règle que `anonymize_csv` — la revue, elle, les
    laisse décochées mais cochables."""
    result = scan_workbook(path, ner, ref)
    retained = {k: [e for e in v if e.confirmed] for k, v in result.scanned.items()}
    retained = {k: v for k, v in retained.items() if v}
    report = apply_workbook(result, retained, ref)
    postprocess_workbook(result.workbook, ner, ref, report)
    purge_metadata(result.workbook, report)
    out = anonymized_path(path, output_dir, when)
    result.workbook.save(out)
    return out, report
