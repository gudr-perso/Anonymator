"""Les exemples « documents » du jeu de démonstration : contrat Word, bulletin
de paie PDF et FEC comptable. On vérifie ce qui les rend utiles en démo —
en-tête/pied de page et liens du .docx, couche texte du .pdf, structure
normative du FEC — et le fait que les règles déterministes y mordent, modèle
GLiNER absent."""
from pathlib import Path

import pytest

from anonymator.files import txt_io
from anonymator.ner import NullNer
from anonymator.pipeline import detect
from anonymator.referential import Referential

EXEMPLES = Path("exemples")
DOCX = EXEMPLES / "Contrat_prestation_Ateliers_Tanguy_EURL.docx"
PAIE = EXEMPLES / "Bulletin_de_paie_2025-06_LACROIX_Damien.pdf"
FEC = EXEMPLES / "404833048FEC20251231.txt"

pytestmark = pytest.mark.skipif(not DOCX.exists(),
                                reason="jeu de démonstration absent")

COLONNES_FEC = [
    "JournalCode", "JournalLib", "EcritureNum", "EcritureDate",
    "CompteNum", "CompteLib", "CompAuxNum", "CompAuxLib",
    "PieceRef", "PieceDate", "EcritureLib", "Debit", "Credit",
    "EcritureLet", "DateLet", "ValidDate", "Montantdevise", "Idevise",
]


def _types_detectes(text: str) -> set[str]:
    ref = Referential.load_default()
    return {e.type for e in detect(text, NullNer(), ref)}


# --------------------------------------------------------------------- contrat
def test_contrat_docx_porte_entete_pied_de_page_et_liens_externes():
    import zipfile
    from docx import Document

    doc = Document(DOCX)
    section = doc.sections[0]
    entete = "\n".join(p.text for p in section.header.paragraphs)
    pied = "\n".join(p.text for p in section.footer.paragraphs)
    assert "ATELIERS TANGUY EURL" in entete
    assert "SIRET 40483304800022" in pied

    with zipfile.ZipFile(DOCX) as z:
        rels = "".join(z.read(n).decode("utf-8") for n in z.namelist()
                       if n.endswith(".xml.rels"))
    # Un lien dans le corps, un autre dans le pied de page.
    assert "https://extranet.ateliers-tanguy.net/login" in rels
    assert "https://www.ateliers-tanguy.net/espace-client" in rels


def test_contrat_docx_porte_des_metadonnees_identifiantes():
    """Ce sont elles que la purge des métadonnées doit retirer."""
    from docx import Document

    core = Document(DOCX).core_properties
    assert core.author == "Delphine Salvatore"
    assert core.last_modified_by == "Sandra Leclerc"


# ---------------------------------------------------------------------- paie
def test_bulletin_de_paie_est_un_pdf_natif_avec_donnees_sensibles():
    from anonymator.files.pdf import extract

    doc = extract.open_document(PAIE)
    try:
        extract.ensure_native(doc)          # lève si scanné
        texte = "\n".join(page.get_text() for page in doc)
    finally:
        doc.close()

    assert "1 84 03 44 109 025 32" in texte            # NIR du salarié
    assert "FR76 1680 6050 1400 0918 2736 413" in texte
    assert {"NIR", "IBAN", "SIRET", "BIRTHDATE", "EMAIL", "PHONE",
            "ADDRESS"} <= _types_detectes(texte)


# ----------------------------------------------------------------------- FEC
def test_fec_respecte_la_structure_normative():
    text, encoding = txt_io.read_text(FEC)
    lignes = text.splitlines()
    assert encoding == "utf-8"
    assert lignes[0].split("\t") == COLONNES_FEC
    assert len(lignes) > 300, "quelques centaines de lignes attendues"
    assert all(len(l.split("\t")) == len(COLONNES_FEC) for l in lignes[1:])


def test_fec_est_equilibre():
    text, _ = txt_io.read_text(FEC)
    debit = credit = 0.0
    for ligne in text.splitlines()[1:]:
        champs = ligne.split("\t")
        debit += float(champs[11].replace(",", "."))
        credit += float(champs[12].replace(",", "."))
    assert round(debit - credit, 2) == 0.0


def test_fec_porte_des_comptes_nominatifs_et_des_libelles_sensibles():
    text, _ = txt_io.read_text(FEC)
    assert "LACROIX Damien - rémunérations dues" in text
    assert "SALVATORE Delphine - compte courant d'associé" in text
    assert {"NIR", "IBAN", "EMAIL", "PHONE", "SIRET",
            "ADDRESS"} <= _types_detectes(text)
