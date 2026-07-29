from datetime import datetime

import openpyxl
from PySide6.QtCore import Qt

from anonymator.referential import Referential
from anonymator.ner import FakeNer
from anonymator.ui.model_loader import ModelLoader
from anonymator.ui.preferences import Preferences
from anonymator.ui.file_screen import FileScreen
from anonymator.core.xlsx_review_session import XlsxReviewSession
from anonymator.core.tabular_review_session import CLEAR, MASK


def _book(path, rows=24):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clients"
    for c, h in enumerate(["code_client", "contact_nom", "secteur", "note"],
                          start=1):
        ws.cell(row=1, column=c, value=h)
    for i in range(rows):
        r = i + 2
        ws.cell(row=r, column=1, value="C%07d" % (i + 1))
        ws.cell(row=r, column=2, value=["Leclerc", "Berger", "Poirier"][i % 3])
        ws.cell(row=r, column=3, value=["Industrie", "BTP", "Textile"][i % 3])
        ws.cell(row=r, column=4, value="note %d" % i)
    ws2 = wb.create_sheet("Tiers")
    ws2["A1"] = "Fournisseur Claire Martin"
    wb.save(path)
    return path


def _screen(tmp_path, mapping=None):
    loader = ModelLoader(FakeNer(mapping or {}))
    return FileScreen(Referential.load_default(), loader,
                      Preferences(output_dir=str(tmp_path)), on_back=lambda: None)


def _reviewed(qtbot, tmp_path, rows=24, mapping=None):
    s = _screen(tmp_path, mapping)
    qtbot.addWidget(s)
    s.load_path(str(_book(tmp_path / "b.xlsx", rows)))
    assert s.btn_review.isEnabled()
    s.analyze()
    qtbot.waitUntil(lambda: s.session is not None, timeout=15000)
    return s


def test_review_is_open_to_xlsx(qtbot, tmp_path):
    s = _reviewed(qtbot, tmp_path)
    assert isinstance(s.session, XlsxReviewSession)
    assert "PERSON" in s.session.types()


def test_grid_shows_the_first_sheet_with_its_headers(qtbot, tmp_path):
    s = _reviewed(qtbot, tmp_path)
    assert s.table.columnCount() == 4
    assert s.table.horizontalHeaderItem(1).text() == "contact_nom"
    assert s.table.item(0, 1).text() == "Leclerc"


def test_sheet_selector_switches_the_grid(qtbot, tmp_path):
    s = _reviewed(qtbot, tmp_path)
    assert [s.sheet_box.itemText(i) for i in range(s.sheet_box.count())] == [
        "Clients", "Tiers"]
    s.sheet_box.setCurrentText("Tiers")
    assert s.table.item(0, 0).text() == "Fournisseur Claire Martin"
    assert s.page == 0


def test_grid_paginates_a_long_sheet(qtbot, tmp_path):
    """_render_units_page ne pagine pas : une feuille de 100 000 lignes gèlerait
    l'interface. La grille, elle, pagine comme le CSV."""
    s = _reviewed(qtbot, tmp_path, rows=45)
    assert s._page_count() == 3
    s._go(99)
    assert s.page == 2
    assert s.table.rowCount() == 5
    s._go(0)
    assert s.table.rowCount() == 20


def _is_highlighted(item):
    """Une cellule sans surlignage garde une brosse NoBrush : c'est le style
    qu'il faut lire, sa couleur restant opaque par défaut."""
    return item.background().style() != Qt.NoBrush


def test_highlighting_follows_the_session(qtbot, tmp_path):
    s = _reviewed(qtbot, tmp_path)
    assert _is_highlighted(s.table.item(0, 1)) is True     # colonne PERSON
    assert _is_highlighted(s.table.item(0, 2)) is False    # nomenclature


def test_column_override_works_on_a_sheet(qtbot, tmp_path):
    s = _reviewed(qtbot, tmp_path)
    s.apply_column_override(3, MASK, "ORG")            # colonne « note »
    assert s.session.column_override(("Clients", 3)) == (MASK, "ORG")
    assert s.session.count_retained("ORG") == 24


def test_apply_writes_the_workbook(qtbot, tmp_path):
    s = _reviewed(qtbot, tmp_path, mapping={"Claire Martin": "PERSON"})
    s.apply_column_override(0, CLEAR)                  # code_client hors périmètre
    res = s.run(when=datetime(2026, 1, 2, 3, 4, 5))
    ws = openpyxl.load_workbook(res.output_path)["Clients"]
    assert ws.cell(row=1, column=2).value == "contact_nom"
    assert ws.cell(row=2, column=1).value == "C0000001"
    assert ws.cell(row=2, column=2).value == "[PERSONNE]"


def test_unchecking_a_value_updates_the_grid(qtbot, tmp_path):
    s = _reviewed(qtbot, tmp_path)
    assert _is_highlighted(s.table.item(0, 1)) is True
    s.session.set_value_enabled("PERSON", "Leclerc", False)
    s._render_current()
    assert _is_highlighted(s.table.item(0, 1)) is False


def test_perimeter_card_stays_hidden_for_xlsx(qtbot, tmp_path):
    s = _reviewed(qtbot, tmp_path)
    assert s.perimetre_card.isHidden()


# ---- interrupteur « première ligne = en-têtes », par feuille ----

def test_header_switch_reflects_the_current_sheet(qtbot, tmp_path):
    s = _reviewed(qtbot, tmp_path)
    assert s.header_switch.isChecked() is True        # Clients a des titres
    s.sheet_box.setCurrentText("Tiers")
    assert s.header_switch.isChecked() is False       # une seule ligne


def test_header_switch_reanalyses_the_sheet(qtbot, tmp_path):
    """Un classeur ne peut pas se re-rendre sans rescan : le choix de
    l'utilisateur ne veut rien dire tant que les colonnes n'ont pas été
    reclassées."""
    from unittest.mock import patch
    from PySide6.QtWidgets import QMessageBox
    s = _reviewed(qtbot, tmp_path)
    with patch("anonymator.ui.file_screen.QMessageBox.question",
               return_value=QMessageBox.Yes):
        s.header_switch.setChecked(False)
    qtbot.waitUntil(lambda: s.session is not None and not s._busy, timeout=15000)
    assert s._xlsx.has_header["Clients"] is False
    assert s._sheet == "Clients"                       # feuille conservée
    assert s.table.item(0, 1).text() == "contact_nom"   # ligne 1 = donnée


def test_header_switch_can_be_declined(qtbot, tmp_path):
    from unittest.mock import patch
    from PySide6.QtWidgets import QMessageBox
    s = _reviewed(qtbot, tmp_path)
    with patch("anonymator.ui.file_screen.QMessageBox.question",
               return_value=QMessageBox.No):
        s.header_switch.setChecked(False)
    assert s.header_switch.isChecked() is True
    assert s._xlsx.has_header["Clients"] is True


def test_header_change_reports_manual_choices(qtbot, tmp_path):
    """Les noms restent détectables sans en-tête (par le NER cette fois, non
    plus par le typage de colonne) : la clé (type, valeur) survit, et le
    décochage avec elle. Le forçage, positionnel, survit aussi."""
    from unittest.mock import patch
    from PySide6.QtWidgets import QMessageBox
    s = _reviewed(qtbot, tmp_path, mapping={"Leclerc": "PERSON",
                                            "Berger": "PERSON",
                                            "Poirier": "PERSON"})
    s.session.set_value_enabled("PERSON", "Berger", False)
    s.apply_column_override(3, MASK, "ORG")
    with patch("anonymator.ui.file_screen.QMessageBox.question",
               return_value=QMessageBox.Yes):
        s.header_switch.setChecked(False)
    qtbot.waitUntil(lambda: s.session is not None and not s._busy, timeout=15000)
    assert s.session.is_value_enabled("PERSON", "Berger") is False
    assert s.session.column_override(("Clients", 3)) == (MASK, "ORG")
