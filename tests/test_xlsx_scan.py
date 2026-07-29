import openpyxl

from anonymator.referential import Referential
from anonymator.ner import FakeNer, NullNer
from anonymator.files import xlsx_io
from anonymator.files.columns import SKIP, TEXT, TYPED


def _book(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clients"
    for c, h in enumerate(["code_client", "contact_nom", "telephone",
                           "secteur", "ca_2025"], start=1):
        ws.cell(row=1, column=c, value=h)
    for i in range(24):
        r = i + 2
        ws.cell(row=r, column=1, value="C%07d" % (i + 1))
        ws.cell(row=r, column=2, value=["Leclerc", "Berger", "Poirier"][i % 3])
        ws.cell(row=r, column=3, value="03 73 41 92 %02d" % i)
        ws.cell(row=r, column=4, value=["Industrie", "BTP", "Textile"][i % 3])
        ws.cell(row=r, column=5, value=1000 + i)
    ws2 = wb.create_sheet("Notes")
    ws2["A1"] = "Fournisseur Claire Martin"
    ws2["A2"] = "=A1"
    wb.save(path)
    return path


def test_scan_reports_sheets_matrices_plans_and_headers(tmp_path):
    src = _book(tmp_path / "b.xlsx")
    res = xlsx_io.scan_workbook(src, NullNer(), Referential.load_default())
    assert res.sheets == ["Clients", "Notes"]
    assert res.has_header["Clients"] is True
    assert res.has_header["Notes"] is False
    assert res.plans["Clients"][1].policy == TYPED
    assert res.plans["Clients"][1].etype == "PERSON"
    assert res.plans["Clients"][4].policy == SKIP       # ca_2025 : mesures
    assert res.matrices["Clients"][0][1] == "contact_nom"


def test_scan_keys_entities_by_sheet_row_and_column(tmp_path):
    src = _book(tmp_path / "b.xlsx")
    res = xlsx_io.scan_workbook(src, NullNer(), Referential.load_default())
    assert ("Clients", 1, 1) in res.scanned            # 1re ligne de données
    assert res.scanned[("Clients", 1, 1)][0].type == "PERSON"
    assert ("Clients", 0, 1) not in res.scanned        # ligne de titres épargnée
    assert not any(k[2] == 4 for k in res.scanned)     # colonne de mesures


def test_scan_never_reads_a_formula(tmp_path):
    src = _book(tmp_path / "b.xlsx")
    res = xlsx_io.scan_workbook(src, FakeNer({"Claire Martin": "PERSON"}),
                                Referential.load_default())
    assert res.matrices["Notes"][1][0] == ""           # =A1 compte pour vide
    assert ("Notes", 1, 0) not in res.scanned


def test_header_override_wins_over_detection(tmp_path):
    """Sans en-tête, la ligne 1 redevient une donnée : le typage par nom de
    colonne tombe, et « telephone » retourne au texte libre."""
    src = _book(tmp_path / "b.xlsx")
    ref = Referential.load_default()
    typed = xlsx_io.scan_workbook(src, NullNer(), ref)
    assert typed.plans["Clients"][2].policy == TYPED
    assert typed.plans["Clients"][2].etype == "PHONE"

    res = xlsx_io.scan_workbook(src, NullNer(), ref,
                                header_overrides={"Clients": False})
    assert res.has_header["Clients"] is False
    assert res.plans["Clients"][2].policy == TEXT
    assert ("Clients", 1, 2) in res.scanned            # données toujours vues


def test_apply_writes_only_the_retained_cells(tmp_path):
    src = _book(tmp_path / "b.xlsx")
    ref = Referential.load_default()
    res = xlsx_io.scan_workbook(src, NullNer(), ref)
    retained = {k: v for k, v in res.scanned.items() if k[2] == 1}   # contact_nom
    report = xlsx_io.apply_workbook(res, retained, ref)
    ws = res.workbook["Clients"]
    assert ws.cell(row=2, column=2).value == "[PERSONNE]"
    assert ws.cell(row=2, column=3).value == "03 73 41 92 00"        # non retenue
    assert any(r["type"] == "PERSON" for r in report.to_rows())
    assert all("Clients!" in r["locations"] for r in report.to_rows())


def test_apply_never_rewrites_a_formula(tmp_path):
    from anonymator.model import Entity
    src = _book(tmp_path / "b.xlsx")
    ref = Referential.load_default()
    res = xlsx_io.scan_workbook(src, NullNer(), ref)
    forged = {("Notes", 1, 0): [Entity("PERSON", "=A1", 0, 3, "column", 1.0)]}
    xlsx_io.apply_workbook(res, forged, ref)
    assert res.workbook["Notes"]["A2"].value == "=A1"
