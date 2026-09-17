"""Fichiers piégés : un témoin unique par zone susceptible de porter une donnée.

Chaque témoin est une adresse e-mail `temoin.<zone>@example.com`. Le choix est
délibéré : l'adresse est le motif que les règles déterministes reconnaissent
sans ambiguïté et sans le modèle. Un témoin qui survit ne dit donc rien de la
qualité de détection, il dit que **la zone n'a pas été lue** — ce qui est
exactement ce que ces tests doivent prouver.

Les fichiers produits servent `test_perimetre_tenu.py`, qui vérifie que ce que
l'interface annonce comme traité l'est réellement.
"""
import re
import zipfile
from pathlib import Path

import fitz
import openpyxl
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from pptx import Presentation
from pptx.util import Inches

TEMOIN_RE = re.compile(rb"temoin\.[a-z0-9-]+@example\.com")


def temoins_in(data: bytes) -> set[str]:
    return {m.decode() for m in TEMOIN_RE.findall(data)}


def temoins_in_archive(path: Path) -> dict[str, list[str]]:
    """{témoin: parties de l'archive où il survit} — vide si tout est masqué."""
    out: dict[str, list[str]] = {}
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            for t in temoins_in(z.read(name)):
                out.setdefault(t, []).append(name)
    return out


# --- Word ---------------------------------------------------------------

def make_docx(path: Path) -> Path:
    """Document Word semé dans toutes les zones annoncées comme traitées."""
    doc = Document()
    doc.add_paragraph("Contact : temoin.corps@example.com")

    # Suppression suivie : le texte retiré vit dans <w:delText>.
    p_del = doc.add_paragraph()
    d = etree.SubElement(p_del._p, qn('w:del'))
    d.set(qn('w:id'), '1'); d.set(qn('w:author'), 'temoin.auteur-revision@example.com')
    d.set(qn('w:date'), '2026-01-01T00:00:00Z')
    r = etree.SubElement(d, qn('w:r'))
    etree.SubElement(r, qn('w:delText')).text = "temoin.suppression@example.com"

    # Insertion suivie.
    p_ins = doc.add_paragraph()
    i = etree.SubElement(p_ins._p, qn('w:ins'))
    i.set(qn('w:id'), '2'); i.set(qn('w:author'), 'Relecteur')
    i.set(qn('w:date'), '2026-01-01T00:00:00Z')
    r2 = etree.SubElement(i, qn('w:r'))
    etree.SubElement(r2, qn('w:t')).text = "temoin.insertion@example.com"

    # Instruction de champ.
    p_f = doc.add_paragraph()
    r_f = etree.SubElement(p_f._p, qn('w:r'))
    etree.SubElement(r_f, qn('w:instrText')).text = \
        ' HYPERLINK "mailto:temoin.champ@example.com" '

    doc.sections[0].header.paragraphs[0].text = "temoin.entete@example.com"
    doc.sections[0].footer.paragraphs[0].text = "temoin.pied@example.com"

    # Lien hypertexte : texte affiché + cible dans les relations.
    rid = doc.part.relate_to(
        "mailto:temoin.cible-lien@example.com",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    p_h = doc.add_paragraph()
    hl = etree.SubElement(p_h._p, qn('w:hyperlink')); hl.set(qn('r:id'), rid)
    r_h = etree.SubElement(hl, qn('w:r'))
    etree.SubElement(r_h, qn('w:t')).text = "temoin.texte-lien@example.com"

    doc.core_properties.author = "temoin.auteur@example.com"
    doc.core_properties.title = "temoin.titre@example.com"
    doc.save(str(path))
    _add_parts(path, {
        "customXml/item1.xml":
            '<?xml version="1.0"?><root><v>temoin.customxml@example.com</v></root>',
    })
    return path


def _add_parts(path: Path, parts: dict[str, str]) -> None:
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        existing = set(zin.namelist())
        for item in zin.infolist():
            zo.writestr(item, zin.read(item.filename))
        for name, content in parts.items():
            if name not in existing:
                zo.writestr(name, content)
    tmp.replace(path)


# --- PowerPoint ---------------------------------------------------------

def make_pptx(path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "temoin.titre-slide@example.com"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    box.text_frame.text = "temoin.zone-texte@example.com"
    slide.notes_slide.notes_text_frame.text = "temoin.notes@example.com"

    layout = slide.slide_layout
    for ph in layout.placeholders:
        ph.text_frame.text = "temoin.disposition@example.com"
        break
    for ph in layout.slide_master.placeholders:
        ph.text_frame.text = "temoin.masque@example.com"
        break

    prs.core_properties.author = "temoin.auteur@example.com"
    prs.core_properties.title = "temoin.titre-doc@example.com"
    prs.save(str(path))
    return path


# --- Classeur -----------------------------------------------------------

def make_xlsx(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Visible"
    ws["A1"] = "contact"
    ws["A2"] = "temoin.cellule@example.com"
    ws["B1"] = "temoin.intitule@example.com"      # identité en ligne de titres
    ws["B2"] = 1
    ws["C2"] = '="temoin.formule@example.com"'
    ws["A2"].comment = Comment("temoin.commentaire@example.com",
                               "temoin.auteur-commentaire@example.com")
    ws.oddHeader.center.text = "temoin.entete@example.com"
    ws.oddFooter.center.text = "temoin.pied@example.com"

    hidden = wb.create_sheet("Masquee")
    hidden["A1"] = "temoin.feuille-masquee@example.com"
    hidden.sheet_state = "hidden"

    # Colonne à faible cardinalité : le profil même d'une nomenclature, et
    # celui d'une colonne d'intervenants.
    axe = wb.create_sheet("Axe")
    axe["A1"] = "intervenant"
    axe["B1"] = "montant"
    for row in range(2, 42):
        axe.cell(row=row, column=1,
                 value=f"temoin.intervenant{(row % 5) + 1}@example.com")
        axe.cell(row=row, column=2, value=row * 10)

    wb.defined_names.add(
        DefinedName("temoin_nom", attr_text='"temoin.nom-defini@example.com"'))
    wb.properties.creator = "temoin.auteur@example.com"
    wb.properties.title = "temoin.titre@example.com"
    wb.save(path)
    return path


def make_csv(path: Path) -> Path:
    """Tableau croisé : identités en intitulés de colonne et dans une colonne
    à faible cardinalité, le tout au milieu de mesures numériques."""
    lines = ["mois;temoin.intitule-a@example.com;temoin.intitule-b@example.com;"
             "intervenant;acte"]
    for i in range(1, 41):
        lines.append(f"2026-{i % 12 + 1:02d};{i * 3};{i * 7};"
                     f"temoin.intervenant{(i % 5) + 1}@example.com;consultation")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


# --- PDF ----------------------------------------------------------------

def make_pdf(path: Path) -> Path:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Contact : temoin.corps@example.com", fontsize=11)
    page.add_text_annot((300, 100), "temoin.annotation@example.com")
    highlight = page.add_highlight_annot(fitz.Rect(72, 150, 300, 170))
    highlight.set_info(content="temoin.surlignage@example.com",
                       title="temoin.auteur-annotation@example.com")
    highlight.update()

    widget = fitz.Widget()
    widget.field_name = "champ1"
    widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
    widget.rect = fitz.Rect(72, 200, 300, 220)
    widget.field_value = "temoin.formulaire@example.com"
    page.add_widget(widget)

    doc.embfile_add("piece.txt", b"temoin.piece-jointe@example.com",
                    filename="piece.txt")
    doc.set_toc([[1, "temoin.signet@example.com", 1]])
    doc.set_metadata({"author": "temoin.auteur@example.com",
                      "title": "temoin.titre@example.com"})
    doc.save(str(path))
    doc.close()
    return path
