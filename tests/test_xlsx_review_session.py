import openpyxl

from anonymator.referential import Referential
from anonymator.ner import FakeNer, NullNer
from anonymator.files import xlsx_io
from anonymator.core.xlsx_review_session import XlsxReviewSession
from anonymator.core.tabular_review_session import CLEAR, MASK


def _book(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clients"
    for c, h in enumerate(["code_client", "contact_nom", "secteur", "note"],
                          start=1):
        ws.cell(row=1, column=c, value=h)
    for i in range(24):
        r = i + 2
        ws.cell(row=r, column=1, value="C%07d" % (i + 1))
        ws.cell(row=r, column=2, value=["Leclerc", "Berger", "Poirier"][i % 3])
        ws.cell(row=r, column=3, value=["Industrie", "BTP", "Textile"][i % 3])
        ws.cell(row=r, column=4, value="note %d" % i)
    ws2 = wb.create_sheet("Tiers")
    ws2["A1"] = "Fournisseur Claire Martin"
    wb.save(path)
    return path


def _session(tmp_path, ner=None):
    src = _book(tmp_path / "b.xlsx")
    ref = Referential.load_default()
    res = xlsx_io.scan_workbook(src, ner or NullNer(), ref)
    return XlsxReviewSession(res, ref), src


def test_types_and_counts(tmp_path):
    s, _src = _session(tmp_path)
    assert "PERSON" in s.types()
    assert s.count_retained("PERSON") == 24        # colonne typée par en-tête


def test_entities_are_keyed_by_sheet(tmp_path):
    s, _src = _session(tmp_path, FakeNer({"Claire Martin": "PERSON"}))
    assert [e.value for e in s.entities_for_cell("Clients", 1, 1)] == ["Leclerc"]
    assert s.entities_for_cell("Tiers", 0, 0)[0].value == "Claire Martin"


def test_column_override_is_per_sheet(tmp_path):
    """Une colonne s'entend « colonne de telle feuille » : forcer une colonne
    de Clients ne doit rien changer sur Tiers."""
    s, _src = _session(tmp_path, FakeNer({"Claire Martin": "PERSON"}))
    s.set_column_override(("Clients", 3), MASK, "ORG")
    assert s.count_retained("ORG") == 24
    assert s.entities_for_cell("Tiers", 0, 0)[0].type == "PERSON"


def test_clear_removes_a_typed_column(tmp_path):
    s, _src = _session(tmp_path)
    s.set_column_override(("Clients", 1), CLEAR)
    assert s.count_retained("PERSON") == 0


def test_apply_and_save_masks_and_preserves_the_rest(tmp_path):
    s, _src = _session(tmp_path, FakeNer({"Claire Martin": "PERSON"}))
    s.set_column_override(("Clients", 3), MASK, "ORG")
    out = tmp_path / "out.xlsx"
    report = s.apply_and_save(out)
    ws = openpyxl.load_workbook(out)["Clients"]
    assert ws.cell(row=1, column=2).value == "contact_nom"   # titres intacts
    assert ws.cell(row=2, column=2).value == "[PERSONNE]"
    assert ws.cell(row=2, column=3).value == "Industrie"     # nomenclature intacte
    assert ws.cell(row=2, column=4).value == "[ORG]"
    assert openpyxl.load_workbook(out)["Tiers"]["A1"].value == "Fournisseur [PERSONNE]"
    assert any(r["locations"].startswith("Clients!") for r in report.to_rows())


def test_unchecked_value_stays_in_clear(tmp_path):
    s, _src = _session(tmp_path)
    s.set_value_enabled("PERSON", "Berger", False)
    out = tmp_path / "out.xlsx"
    s.apply_and_save(out)
    ws = openpyxl.load_workbook(out)["Clients"]
    assert ws.cell(row=3, column=2).value == "Berger"        # décochée
    assert ws.cell(row=2, column=2).value == "[PERSONNE]"


def test_has_column_guards_a_replayed_override(tmp_path):
    s, _src = _session(tmp_path)
    assert s.has_column(("Clients", 3)) is True
    assert s.has_column(("Clients", 99)) is False
    assert s.has_column(("Absente", 0)) is False
